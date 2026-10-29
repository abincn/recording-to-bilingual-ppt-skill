#!/usr/bin/env python3
"""Validate the generated PPTX, cue mapping, subtitles, colors and media durations."""
from __future__ import annotations
import argparse
import re
from pathlib import Path
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pipeline_core import check_cues, palette, read_json, require_translation, rgb, safe_project_file
from make_ass import slide_events
from make_videos import duration

def validate(manifest, ppt):
    manifest, ppt = Path(manifest).resolve(), Path(ppt).resolve()
    data = read_json(manifest)
    folder = manifest.parent
    cues = check_cues(read_json(folder/"merged_cues.json"))
    cn = require_translation(cues, read_json(folder/"sentence_cn.json"))
    specs = read_json(folder/"slides.json")
    prs = Presentation(ppt)
    if len(prs.slides) != len(data["slides"]):
        raise ValueError("PPT slide count differs from manifest")
    mode = data.get("mode", "image")
    expected_cues = []
    for index, (spec, item, slide) in enumerate(zip(
            (specs[str(i)] for i in range(1, len(prs.slides)+1)), data["slides"], prs.slides), 1):
        slide_cues = slide_events(cues, spec)
        if not spec.get("review", False):
            expected_cues.extend(idx for _, _, idx in slide_cues)
        if len([s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.MEDIA]) != 1:
            raise ValueError(f"Slide {index}: expected exactly one embedded video")
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if mode == "image" and len(pictures) != 1:
            raise ValueError(f"Slide {index}: expected one slide picture")
        if mode == "editable" and pictures:
            raise ValueError(f"Slide {index}: editable mode must have no flattened slide picture")
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError(f"Slide {index}: notes missing")
        background = item["bar_color"]
        palette(background, spec.get("text_color"))
        ass = safe_project_file(folder, item["ass"])
        dialogues = [line for line in ass.read_text(encoding="utf-8").splitlines()
                     if line.startswith("Dialogue:")]
        if len(dialogues) != len(slide_cues):
            raise ValueError(f"Slide {index}: ASS cue count does not match timeline")
        for text in dialogues:
            if not re.search(r"[\u3400-\u9fff]", text.split(",", 9)[-1]):
                raise ValueError(f"Slide {index}: dialogue lacks Chinese")
        vid = safe_project_file(folder, item["video"])
        aud = safe_project_file(folder, item["audio"])
        delta = abs(duration(vid)-duration(aud))
        if delta > .25:
            raise ValueError(f"Slide {index}: video duration differs from audio by {delta:.2f}s")
        png = folder/"qa"/f"slide{index:02d}.png"
        if not png.exists():
            raise ValueError(f"Slide {index}: missing video sample frame {png}")
        for sample in sorted(png.parent.glob(png.stem+"*.png")):
            with Image.open(sample) as original:
                frame = original.convert("RGB")
            reference = Image.new("RGB", frame.size, rgb(background))
            mask = ImageChops.difference(frame, reference).convert("L").point(lambda p: 255 if p > 24 else 0)
            bbox = mask.getbbox()
            if bbox is None:
                raise ValueError(f"Slide {index}: no subtitle text in {sample}")
            if bbox[1] < 7 or bbox[3] > frame.height-7:
                raise ValueError(f"Slide {index}: subtitle in {sample} touches safety margin")
    if sorted(expected_cues) != list(range(1, len(cues)+1)):
        raise ValueError("Every transcript cue must appear exactly once across normal slides")
    return {"slides":len(prs.slides), "cues":len(cues), "mode":mode, "ppt_bytes":ppt.stat().st_size}

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--ppt", required=True)
    args = ap.parse_args()
    print(validate(args.manifest, args.ppt))

if __name__ == "__main__":
    main()
