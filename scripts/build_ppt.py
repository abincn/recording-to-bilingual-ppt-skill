#!/usr/bin/env python3
"""Build a clean PPTX containing per-slide original-video subtitles and optional editable text."""
from __future__ import annotations
import argparse
import math
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt
from pipeline_core import palette, read_json, safe_project_file, rgb

def set_autoplay(slide):
    """python-pptx movie timing is indefinite (click); 0 starts at slide entry."""
    videos = slide._element.xpath(".//p:video")
    if len(videos) != 1:
        raise ValueError("Expected exactly one video timing element on this slide")
    starts = videos[0].xpath(".//p:stCondLst/p:cond")
    if len(starts) != 1:
        raise ValueError("Cannot locate video start condition")
    starts[0].set("delay", "0")

def add_editable_text(slide, title, bullets, width, height, bg):
    # All visible slide words are PowerPoint text, never a flattened image.
    dark = RGBColor(*rgb(palette(bg)["text_color"]))
    margin = Inches(.7)
    title_box = slide.shapes.add_textbox(margin, Inches(.35), width-2*margin, Inches(1.05))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.clear()
    para = tf.paragraphs[0]
    para.text = str(title)
    para.font.name = "Noto Sans CJK SC"
    para.font.bold = True
    para.font.size = Pt(30)
    para.font.color.rgb = dark
    para.alignment = PP_ALIGN.LEFT
    body_height = height-Inches(1.9)
    if body_height < Inches(.5):
        raise ValueError("Too little slide area above the subtitle strip")
    bullets = list(bullets)[:6]
    available_pt = body_height / 12700
    content_width_pt = (width-2*margin) / 12700
    body_size = None
    for candidate in (20, 18, 16, 14):
        approx_chars_per_line = max(8, content_width_pt / (candidate*.59))
        rows = sum(max(1, math.ceil(len(str(item))/approx_chars_per_line)) for item in bullets)
        if rows*candidate*1.30 + len(bullets)*8 <= available_pt:
            body_size = candidate
            break
    if body_size is None:
        raise ValueError("Editable text would overflow into the subtitle bar; reduce content")
    body = slide.shapes.add_textbox(margin, Inches(1.6), width-2*margin, body_height)
    body_tf = body.text_frame
    body_tf.word_wrap = True
    body_tf.vertical_anchor = MSO_ANCHOR.TOP
    body_tf.clear()
    for i, item in enumerate(bullets):
        p = body_tf.paragraphs[0] if i == 0 else body_tf.add_paragraph()
        p.text = "• " + str(item)
        p.font.name = "Noto Sans CJK SC"
        p.font.size = Pt(body_size)
        p.font.color.rgb = dark
        p.space_after = Pt(8)
    if len(str(title)) > 130 or any(len(str(b)) > 200 for b in bullets):
        raise ValueError("Slide text too long for a readable editable layout")


def add_editable_table(slide, item, width, height, bg):
    """Render any discipline's structured data as a native PowerPoint table."""
    table_info = item["table"]
    headers, rows = table_info["columns"], table_info["rows"]
    ink = RGBColor(*rgb(palette(bg)["text_color"]))
    light = RGBColor(235, 244, 247)
    strong = RGBColor(205, 227, 233)
    white = RGBColor(255, 255, 255)
    title = slide.shapes.add_textbox(Inches(.65), Inches(.30),
                                     width-Inches(1.3), Inches(.86))
    p = title.text_frame.paragraphs[0]
    p.text = item["title"]
    p.font.name = "Noto Sans CJK SC"
    p.font.size = Pt(29)
    p.font.bold = True
    p.font.color.rgb = ink
    footnote = table_info.get("footnote", "")
    table_height = height-Inches(1.65 if not footnote else 2.12)
    if table_height <= 0:
        raise ValueError("Not enough space above subtitle strip for editable table")
    frame = slide.shapes.add_table(
        len(rows)+1, len(headers), Inches(.65), Inches(1.28),
        width-Inches(1.3), table_height)
    table = frame.table
    for ri, row in enumerate([{"cells":headers,"emphasis":True}]+rows):
        for ci, value in enumerate(row["cells"]):
            cell = table.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = strong if ri == 0 else (light if row["emphasis"] else white)
            cell.margin_left = Inches(.10)
            cell.margin_right = Inches(.07)
            cell.margin_top = Inches(.04)
            cell.margin_bottom = Inches(.03)
            cell.text_frame.word_wrap = True
            cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = str(value)
            p.font.name = "Noto Sans CJK SC"
            p.font.size = Pt(15 if len(headers) <= 4 else 13)
            p.font.bold = ri == 0 or row["emphasis"]
            p.font.color.rgb = ink
    if footnote:
        box = slide.shapes.add_textbox(Inches(.65), height-Inches(.72),
                                       width-Inches(1.3), Inches(.56))
        p = box.text_frame.paragraphs[0]
        p.text = "Note / 注：" + footnote
        p.font.name = "Noto Sans CJK SC"
        p.font.size = Pt(12)
        p.font.color.rgb = ink

def build(manifest, output, *, mode=None, autoplay=None, overwrite=False):
    manifest = Path(manifest).resolve()
    data = read_json(manifest)
    if not data.get("slides"):
        raise ValueError("Manifest requires nonempty slides")
    output = Path(output).resolve()
    if output.exists() and not overwrite:
        raise FileExistsError(f"Output exists (use --force to overwrite): {output}")
    deck = Presentation()
    deck.slide_width = Inches(float(data.get("slide_width_in", 13.333)))
    deck.slide_height = Inches(float(data.get("slide_height_in", 7.5)))
    height = deck.slide_height
    width = deck.slide_width
    strip = int(height * float(data.get("subtitle_fraction", 1.805/7.5)))
    if not 0.15 <= strip/height <= 0.40:
        raise ValueError("Subtitle strip must occupy between 15% and 40% of slide height")
    top_height = height-strip
    blank = deck.slide_layouts[6]
    mode = mode or data.get("mode", "image")
    autoplay = bool(data.get("autoplay", False)) if autoplay is None else autoplay
    if mode not in {"image", "editable"}:
        raise ValueError("Mode must be image or editable")
    for i, item in enumerate(data["slides"], 1):
        colors = palette(item.get("bar_color", data.get("bar_color")), item.get("text_color"))
        slide = deck.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(*rgb(colors["bar_color"]))
        if mode == "image":
            from PIL import Image
            path = safe_project_file(manifest.parent, item["picture"])
            with Image.open(path) as img:
                iw, ih = img.size
            fitted_width = min(int(width), int(top_height * iw/ih))
            fitted_height = min(int(top_height), int(width * ih/iw))
            slide.shapes.add_picture(str(path), int((width-fitted_width)/2), int((top_height-fitted_height)/2),
                                     width=fitted_width, height=fitted_height)
        else:
            if item.get("layout") in {"table", "comparison"}:
                add_editable_table(slide, item, width, top_height, colors["bar_color"])
            else:
                add_editable_text(slide, item["title"], item.get("bullets", []),
                                  width, top_height, colors["bar_color"])
        slide.shapes.add_movie(
            str(safe_project_file(manifest.parent, item["video"])),
            0, top_height, width, strip,
            poster_frame_image=str(safe_project_file(manifest.parent, item["poster"])),
            mime_type="video/mp4",
        )
        if autoplay:
            set_autoplay(slide)
        frame = slide.notes_slide.notes_text_frame
        if frame is not None:
            frame.text = str(item.get("notes", f"Slide {i}: cleaned original audio; bilingual centred subtitles"))
    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    return output

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--manifest", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--mode", choices=["image", "editable"])
    ap.add_argument("--autoplay", action="store_true", default=None)
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()
    path = build(args.manifest, args.out, mode=args.mode, autoplay=args.autoplay, overwrite=args.force)
    print(f"Saved {path} ({path.stat().st_size} bytes)")

if __name__ == "__main__":
    main()
