"""Orchestrator smoke test with stub ASR/translation; no model or API dependency."""
import sys
import tempfile
import subprocess
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]/"scripts"))
import run_pipeline
from validate import validate
from pipeline_core import write_json, read_json
from pptx import Presentation

class PipelineTests(unittest.TestCase):
    def test_end_to_end_image_then_editable_resume(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root/"source.wav"
            subprocess.run(["ffmpeg","-loglevel","error","-f","lavfi","-i",
                            "sine=frequency=440:duration=3.5","-y",str(source)],check=True)
            project = root/"generated"
            cues = [{"start":.2,"end":.8,"en":"Revenue increased."},
                    {"start":1.8,"end":2.5,"en":"Check cash flow."}]
            cn = {1:"收入增加。", 2:"检查现金流量。"}
            first = ["run_pipeline.py", "--audio",str(source),"--project",str(project),
                     "--slides","2","--colors","00FF00,F6FBF3","--mode","image",
                     "--translator","ollama","--autoplay"]
            with patch.object(sys,"argv",first), patch.object(run_pipeline,"transcribe",return_value=cues), \
                 patch.object(run_pipeline,"translate",return_value=cn):
                run_pipeline.main()
            self.assertEqual(validate(project/"manifest.json",project/"courseware-image.pptx")["slides"],2)
            second = ["run_pipeline.py","--audio",str(source),"--project",str(project),
                      "--slides","2","--colors","00FF00,F6FBF3","--mode","editable",
                      "--translations",str(project/"sentence_cn.json"),"--resume","--autoplay"]
            with patch.object(sys,"argv",second):
                run_pipeline.main()
            self.assertEqual(validate(project/"manifest.json",project/"courseware-editable.pptx")["mode"],"editable")

    def test_subject_agnostic_curated_plan_uses_native_table(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root/"voice.wav"
            subprocess.run(["ffmpeg","-loglevel","error","-f","lavfi","-i",
                            "sine=frequency=440:duration=3.2","-y",str(source)],check=True)
            project = root/"lesson"
            project.mkdir()
            cues = [{"start":.2,"end":.8,"en":"Plants capture sunlight."},
                    {"start":1.6,"end":2.5,"en":"They produce oxygen."}]
            write_json(project/"merged_cues.json", cues)
            write_json(project/"sentence_cn.json", {"1":"植物吸收阳光。","2":"它们产生氧气。"})
            write_json(project/"plan.json", {"schema_version":1,"slides":[
                {"cue_start":1,"cue_end":1,"title":"Energy | 能量",
                 "layout":"steps","bullets":["Sunlight / 阳光"]},
                {"cue_start":2,"cue_end":2,"title":"Outputs | 产物",
                 "layout":"table","table":{"columns":["Material / 物质","Output / 产物"],
                                           "rows":[["Air / 空气","Oxygen / 氧气"]]}}
            ]})
            common = ["run_pipeline.py","--audio",str(source),"--project",str(project),
                      "--resume","--translations",str(project/"sentence_cn.json"),
                      "--courseware-plan",str(project/"plan.json"),"--autoplay"]
            for mode in ("image","editable"):
                with patch.object(sys,"argv",common+["--mode",mode]):
                    run_pipeline.main()
                self.assertEqual(validate(project/"manifest.json",
                                          project/f"courseware-{mode}.pptx")["slides"], 2)
            self.assertFalse(read_json(project/"manifest.json")["draft"])
            slides = Presentation(project/"courseware-editable.pptx").slides
            self.assertTrue(any(shape.has_table for shape in slides[1].shapes))

if __name__ == "__main__":
    unittest.main()
